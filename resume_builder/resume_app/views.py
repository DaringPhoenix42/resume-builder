from docx import Document
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from docx import Document
from io import BytesIO

from .models import CoverLetter
from .forms import CoverLetterForm
from django.shortcuts import render, redirect

from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from docx import Document
from docx.shared import Pt
import datetime

from .models import CoverLetter
from .forms import CoverLetterForm
from django.shortcuts import get_object_or_404
from django.template.loader import render_to_string
from django.http import HttpResponse
import pdfkit
from io import BytesIO
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from django.shortcuts import get_object_or_404
from django.http import HttpResponse
from django.template.loader import render_to_string
from django.contrib.admin.views.decorators import staff_member_required

import pdfkit
from io import BytesIO
from docx import Document
from docx.shared import Pt
from docx.enum.style import WD_STYLE_TYPE

from .models import Resume

from django.shortcuts import get_object_or_404
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
import datetime

# For DOCX
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# For PDF
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from .models import CoverLetter
from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib.admin.views.decorators import staff_member_required
from django.contrib.auth.models import User

from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from docx import Document
from io import BytesIO

from .models import Resume, Resource
from .forms import ResumeForm, ResourceForm

def index(request):
    return render(request, 'index.html')

def resume_form(request):
    template_id = request.GET.get('template_id', 1)  # Default to template 1 if none provided
    if request.method == 'POST':
        form = ResumeForm(request.POST)
        if form.is_valid():
            resume = form.save(commit=False)
            resume.template_id = template_id
            # Optionally set resume.user = request.user if you want
            resume.save()
            return redirect('resume_preview', id=resume.id)
    else:
        form = ResumeForm(initial={'template_id': template_id})
    return render(request, 'resume_form.html', {'form': form, 'template_id': template_id})

def resume_preview(request, id):
    resume = get_object_or_404(Resume, id=id)
    return render(request, 'resume_preview.html', {'resume': resume})

def about(request):
    return render(request, 'about.html')

def contact(request):
    return render(request, 'contact.html')

def templates(request):
    return render(request, 'templates.html')

def faq(request):
    return render(request, 'faq.html')

def pricing(request):
    return render(request, 'pricing.html')

def resources(request):
    return render(request, 'resources.html')

@login_required
def dashboard(request):
    user_resumes = Resume.objects.filter(user=request.user)
    return render(request, 'dashboard.html', {'resumes': user_resumes})

def login_user(request):
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')
        user = authenticate(request, username=username, password=password)
        if user:
            login(request, user)
            return redirect('dashboard')
        else:
            # Invalid credentials
            return render(request, 'login.html', {'error': 'Invalid username or password'})
    return render(request, 'login.html')

def logout_user(request):
    logout(request)
    return redirect('index')

def testimonials(request):
    return render(request, 'testimonials.html')

def examples(request):
    return render(request, 'examples.html')

def blog(request):
    return render(request, 'blog.html')

def cover_letter(request):
    return render(request, 'cover_letter.html')

@staff_member_required
def create_resource(request):
    if request.method == 'POST':
        form = ResourceForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('resources')  # or any other desired page
    else:
        form = ResourceForm()
    return render(request, 'create_resource.html', {'form': form})




@staff_member_required
def create_resource(request):
    # Your existing create_resource logic here (if any)
    pass





def download_resume(request, id, format):
    """
    Generates a PDF or Word resume for the given Resume object.
    """
    resume = get_object_or_404(Resume, id=id)

    if format == 'pdf':
        # 1) Render your HTML template with the resume data
        html_string = render_to_string('resume_pdf_template.html', {'resume': resume})

        # 2) Convert the HTML to PDF using pdfkit
        pdf_file = pdfkit.from_string(html_string, False)  # returns PDF bytes

        # 3) Return as a downloadable PDF
        response = HttpResponse(pdf_file, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{resume.name}_resume.pdf"'
        return response

    elif format == 'word':
        document = Document()

        # --- Centered Name ---
        name_heading = document.add_heading(resume.name or "Your Name", 0)
        name_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # --- Centered Job Title ---
        if resume.job_title:
            job_title_par = document.add_paragraph(resume.job_title)
            job_title_par.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # --- Vertical Contact Info, also centered ---
        contact_par = document.add_paragraph()
        contact_par.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if resume.address:
            contact_par.add_run(f"{resume.address}\n")
        if resume.phone:
            contact_par.add_run(f"{resume.phone}\n")
        if resume.email:
            contact_par.add_run(f"{resume.email}\n")
        if resume.github:
            contact_par.add_run(f"GitHub: {resume.github}\n")
        if resume.linkedin:
            contact_par.add_run(f"LinkedIn: {resume.linkedin}\n")
        if resume.portfolio:
            contact_par.add_run(f"Portfolio: {resume.portfolio}\n")

        # --- About Me ---
        if resume.summary:
            document.add_heading("About Me", level=1)
            document.add_paragraph(resume.summary)

        # --- Education ---
        if resume.educations.exists():
            document.add_heading("Education", level=1)
            for edu in resume.educations.all():
                edu_par = document.add_paragraph(style='List Bullet')
                line = f"{edu.institution}"
                if edu.degree:
                    line += f" - {edu.degree}"
                if edu.field_of_study:
                    line += f" ({edu.field_of_study})"
                if edu.start_year or edu.end_year:
                    line += f", {edu.start_year} - {edu.end_year}"
                edu_par.add_run(line)
                if edu.description:
                    edu_par.add_run(f"\n{edu.description}")

        # --- Experience ---
        if resume.experience:
            document.add_heading("Experience", level=1)
            exp_lines = resume.experience.split('\n')
            for line in exp_lines:
                line = line.strip()
                if line:
                    document.add_paragraph(line, style='List Bullet')

        # --- Skills ---
        if resume.skills:
            document.add_heading("Skills", level=1)
            skill_lines = resume.skills.split('\n')
            for line in skill_lines:
                line = line.strip()
                if line:
                    document.add_paragraph(line, style='List Bullet')

        # --- Certifications ---
        if resume.certifications:
            document.add_heading("Certifications", level=1)
            cert_lines = resume.certifications.split('\n')
            for line in cert_lines:
                line = line.strip()
                if line:
                    document.add_paragraph(line, style='List Bullet')

        # --- Languages ---
        if resume.languages:
            document.add_heading("Languages", level=1)
            lang_lines = resume.languages.split('\n')
            for line in lang_lines:
                line = line.strip()
                if line:
                    document.add_paragraph(line, style='List Bullet')

        # --- Hobbies & Interests ---
        if resume.hobbies_interests:
            document.add_heading("Hobbies & Interests", level=1)
            hobby_lines = resume.hobbies_interests.split('\n')
            for line in hobby_lines:
                line = line.strip()
                if line:
                    document.add_paragraph(line, style='List Bullet')

        # --- Return as a .docx file ---
        buffer = BytesIO()
        document.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer,
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f'attachment; filename="{resume.name}_resume.docx"'
        return response

    else:
        return HttpResponse("Invalid format", status=400)


from django.contrib.auth.forms import UserCreationForm

def signup(request):
    if request.method == 'POST':
        form = UserCreationForm(request.POST)
        if form.is_valid():
            user = form.save()
            # Optionally log the user in right away
            login(request, user)
            return redirect('dashboard')
    else:
        form = UserCreationForm()
    return render(request, 'signup.html', {'form': form})

def login_user(request):
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')
        user = authenticate(request, username=username, password=password)
        if user:
            login(request, user)
            return redirect('dashboard')
        else:
            return render(request, 'login.html', {'error': 'Invalid username or password'})
    return render(request, 'login.html')

def logout_user(request):
    logout(request)
    return redirect('index')




@login_required
def cover_letter(request):
    """Create or edit a cover letter."""
    if request.method == 'POST':
        form = CoverLetterForm(request.POST)
        if form.is_valid():
            cl = form.save(commit=False)
            cl.user = request.user
            cl.save()
            return redirect('cover_letter_preview', id=cl.id)
    else:
        form = CoverLetterForm()
    return render(request, 'cover_letter.html', {'form': form})

@login_required
def cover_letter_preview(request, id):
    """Display a styled preview of the cover letter."""
    cover_letter = get_object_or_404(CoverLetter, id=id, user=request.user)
    return render(request, 'cover_letter_preview.html', {'cover_letter': cover_letter})


@login_required
def download_cover_letter(request, id, format):
    cover_letter = get_object_or_404(CoverLetter, id=id, user=request.user)

    # Build up a "structured" version of the letter data
    today_str = datetime.date.today().strftime("%B %d, %Y")  # e.g. "February 27, 2025"

    # We'll pass the entire cover_letter object + some extra fields to our generator
    # so it can do the layout. 
    letter_context = {
        'full_name': cover_letter.full_name or "Your Name",
        'job_title': cover_letter.job_title or "",
        'phone': cover_letter.phone or "",
        'city_state': cover_letter.city_state or "",
        'email': cover_letter.email or "",
        'linkedin': cover_letter.linkedin or "",
        'date_str': today_str,

        'employer_name': cover_letter.employer_name or "",
        'employer_title': cover_letter.employer_title or "",
        'employer_company': cover_letter.employer_company or "",
        'employer_email': cover_letter.employer_email or "",
        'employer_location': cover_letter.employer_location or "",

        'greeting': cover_letter.greeting or "Dear Hiring Manager,",
        'body': cover_letter.body or "",
        'closing': cover_letter.closing or "Thank you for your time,\nI look forward to hearing from you."
    }

    if format == 'docx':
        return _generate_cover_letter_docx(letter_context)
    elif format == 'pdf':
        return _generate_cover_letter_pdf(letter_context)
    else:
        return HttpResponse("Invalid format", status=400)



from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from django.http import HttpResponse

def _generate_cover_letter_docx(letter_context):
    """
    Creates a .docx file that looks more like your reference image:
    - White background page
    - Dark-blue header bar with white text for name/job/contact
    - Black text in the main body
    """
    document = Document()

    # ---------------------------
    # 1) Set Page Margins & Base Font
    # ---------------------------
    # Adjust margins to ~1 inch all around
    for section in document.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Set "Normal" style to black text, Calibri 11pt
    style = document.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0, 0, 0)  # black text

    # ---------------------------
    # 2) Create Header Table
    # ---------------------------
    # Single row, two columns: left (name/job), right (contact)
    header_table = document.add_table(rows=1, cols=2)
    header_table.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    header_table.autofit = False
    # Column widths
    header_table.columns[0].width = Inches(3.5)
    header_table.columns[1].width = Inches(2.5)

    # Shading color: a dark blue (e.g. #34495E)
    # Convert hex color to fill attribute
    shading_xml = r'<w:shd {} w:fill="34495E" w:color="auto" w:val="clear"/>'.format(nsdecls('w'))

    row = header_table.rows[0]
    for cell in row.cells:
        cell_properties = cell._tc.get_or_add_tcPr()
        cell_properties.append(parse_xml(shading_xml))

    # Left cell: Name (big/bold/white) + Job Title (white)
    left_cell = row.cells[0]
    left_par = left_cell.paragraphs[0]

    run_name = left_par.add_run(letter_context['full_name'])
    run_name.font.size = Pt(16)
    run_name.font.bold = True
    run_name.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)  # white

    left_par.add_run("\n")
    run_title = left_par.add_run(letter_context['job_title'])
    run_title.font.size = Pt(12)
    run_title.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)  # white

    # Right cell: contact info (white text, joined by |)
    right_cell = row.cells[1]
    right_par = right_cell.paragraphs[0]
    contact_parts = []
    if letter_context['phone']:
        contact_parts.append(letter_context['phone'])
    if letter_context['city_state']:
        contact_parts.append(letter_context['city_state'])
    if letter_context['email']:
        contact_parts.append(letter_context['email'])
    if letter_context['linkedin']:
        contact_parts.append(letter_context['linkedin'])

    contact_str = " | ".join(contact_parts)
    run_contact = right_par.add_run(contact_str)
    run_contact.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)  # white

    # Add spacing after the header table
    document.add_paragraph()  # blank paragraph

    # ---------------------------
    # 3) Date & Employer Info
    # ---------------------------
    # Example: "February 27, 2025"
    document.add_paragraph(letter_context['date_str'])

    # If you have multiple lines for the employer:
    if letter_context['employer_name'] or letter_context['employer_title']:
        line = f"{letter_context['employer_name']}, {letter_context['employer_title']}"
        document.add_paragraph(line)
    if letter_context['employer_company']:
        document.add_paragraph(letter_context['employer_company'])
    if letter_context['employer_email']:
        document.add_paragraph(letter_context['employer_email'])
    if letter_context['employer_location']:
        document.add_paragraph(letter_context['employer_location'])

    document.add_paragraph()  # extra blank line

    # ---------------------------
    # 4) Greeting, Body, Closing
    # ---------------------------
    # Greeting in bold
    greeting_par = document.add_paragraph(letter_context['greeting'])
    greeting_par.runs[0].font.bold = True

    # Body: Split on double newlines if you want multiple paragraphs
    body_blocks = letter_context['body'].split('\n\n')
    for block in body_blocks:
        p = document.add_paragraph(block)
        p.paragraph_format.space_after = Pt(10)  # spacing after paragraphs

    # Closing
    document.add_paragraph(letter_context['closing'])
    # Signature line
    document.add_paragraph("Sincerely,\n" + letter_context['full_name'])

    # ---------------------------
    # 5) Return as .docx
    # ---------------------------
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    response['Content-Disposition'] = 'attachment; filename="cover_letter.docx"'
    document.save(response)
    return response





from django.http import HttpResponse
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import inch

def _generate_cover_letter_pdf(letter_context):
    """
    Creates a PDF with a white page background and a single dark-blue header row
    that has white text for name/job/contact. The rest is black text.
    """
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="cover_letter.pdf"'

    doc = SimpleDocTemplate(
        response,
        pagesize=letter,
        leftMargin=72, rightMargin=72, topMargin=72, bottomMargin=72
    )

    styles = getSampleStyleSheet()
    normal_style = ParagraphStyle(
        'NormalCustom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=11,
        leading=15,
        textColor=colors.black,  # black text
    )
    bold_style = ParagraphStyle(
        'BoldCustom',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.black,
    )

    story = []

    # 1) Dark-blue header row with white text
    left_html = f"<font color='white'><b>{letter_context['full_name']}</b><br/>{letter_context['job_title']}</font>"
    contact_parts = []
    if letter_context['phone']:
        contact_parts.append(letter_context['phone'])
    if letter_context['city_state']:
        contact_parts.append(letter_context['city_state'])
    if letter_context['email']:
        contact_parts.append(letter_context['email'])
    if letter_context['linkedin']:
        contact_parts.append(letter_context['linkedin'])
    right_html = f"<font color='white'>{' | '.join(contact_parts)}</font>"

    header_table_data = [
        [Paragraph(left_html, normal_style), Paragraph(right_html, normal_style)]
    ]
    header_table = Table(header_table_data, colWidths=[3.5*inch, 2.5*inch])
    header_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#34495E')),  # dark-blue
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('LEFTPADDING', (0,0), (-1,-1), 10),
        ('RIGHTPADDING', (0,0), (-1,-1), 10),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
    ]))
    story.append(header_table)
    story.append(Spacer(1, 12))

    # 2) Date & Employer Info
    story.append(Paragraph(letter_context['date_str'], normal_style))
    if letter_context['employer_name'] or letter_context['employer_title']:
        line = f"{letter_context['employer_name']}, {letter_context['employer_title']}"
        story.append(Paragraph(line, normal_style))
    if letter_context['employer_company']:
        story.append(Paragraph(letter_context['employer_company'], normal_style))
    if letter_context['employer_email']:
        story.append(Paragraph(letter_context['employer_email'], normal_style))
    if letter_context['employer_location']:
        story.append(Paragraph(letter_context['employer_location'], normal_style))
    story.append(Spacer(1, 12))

    # 3) Greeting
    story.append(Paragraph(letter_context['greeting'], bold_style))
    story.append(Spacer(1, 10))

    # 4) Body
    paragraphs = letter_context['body'].split('\n\n')
    for block in paragraphs:
        block_html = block.replace('\n', '<br/>')
        story.append(Paragraph(block_html, normal_style))
        story.append(Spacer(1, 10))

    # 5) Closing + signature
    closing_html = letter_context['closing'].replace('\n', '<br/>')
    story.append(Paragraph(closing_html, normal_style))
    story.append(Spacer(1, 12))
    story.append(Paragraph("Sincerely,<br/>" + letter_context['full_name'], normal_style))

    doc.build(story)
    return response






from django.shortcuts import render, redirect, get_object_or_404
from .models import Resume
from .forms import ResumeForm, EducationForm, EducationFormSet

def resume_form_view(request, resume_id=None):
    if resume_id:
        resume = get_object_or_404(Resume, pk=resume_id)
    else:
        resume = Resume()

    if request.method == 'POST':
        form = ResumeForm(request.POST, instance=resume)
        formset = EducationFormSet(request.POST, instance=resume)
        if form.is_valid() and formset.is_valid():
            resume = form.save()
            formset.save()
            return redirect('resume_preview', id=resume.id)
    else:
        form = ResumeForm(instance=resume)
        formset = EducationFormSet(instance=resume)

    return render(request, 'resume_form.html', {
        'form': form,
        'formset': formset
    })
